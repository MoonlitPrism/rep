import sys
import csv
import os


# Необходимо подгрузить PySide6 openpyxl python-docx

from PySide6.QtCore import Qt, QAbstractTableModel
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLabel, QPushButton, QLineEdit,
    QListWidget, QListWidgetItem, QStackedWidget, QVBoxLayout, QHBoxLayout,
    QFileDialog, QMessageBox, QTableView, QTextEdit, QHeaderView
)


class CsvTableModel(QAbstractTableModel):
    def __init__(self, rows=None, headers=None):
        super().__init__()
        self._rows = rows or []
        self._headers = headers or []

    def rowCount(self, parent=None):
        return len(self._rows)

    def columnCount(self, parent=None):
        if self._headers:
            return len(self._headers)
        return len(self._rows[0]) if self._rows else 0

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        if role in (Qt.DisplayRole, Qt.EditRole):
            try:
                return str(self._rows[index.row()][index.column()])
            except IndexError:
                return ""
        return None

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        if orientation == Qt.Horizontal and self._headers:
            if section < len(self._headers):
                return str(self._headers[section])
            return ""
        return str(section + 1)

    def set_data(self, rows, headers):
        self.beginResetModel()
        self._rows = rows
        self._headers = headers
        self.endResetModel()


class MainScreen(QWidget):
    def __init__(self, on_select_file, on_open_recent):
        super().__init__()
        self._on_select_file = on_select_file
        self._on_open_recent = on_open_recent

        layout = QVBoxLayout(self)

        title = QLabel("Просмотрщик файлов")
        title.setAlignment(Qt.AlignCenter)
        font = title.font()
        font.setPointSize(18)
        font.setBold(True)
        title.setFont(font)
        layout.addWidget(title)

        subtitle = QLabel("Поддерживаемые форматы: .csv, .xlsx, .txt, .docx")
        subtitle.setAlignment(Qt.AlignCenter)
        layout.addWidget(subtitle)

        layout.addSpacing(10)

        self.path_label = QLabel("Файл не выбран")
        self.path_label.setWordWrap(True)
        self.path_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(self.path_label)

        self.select_btn = QPushButton("Выбрать файл")
        self.select_btn.clicked.connect(self._on_select_file)
        layout.addWidget(self.select_btn)

        layout.addSpacing(15)

        recent_label = QLabel("Недавние файлы:")
        layout.addWidget(recent_label)

        self.recent_list = QListWidget()
        self.recent_list.itemDoubleClicked.connect(self._on_recent_double_clicked)
        layout.addWidget(self.recent_list)

        layout.addStretch()

    def _on_recent_double_clicked(self, item: QListWidgetItem):
        path = item.data(Qt.UserRole)
        if path:
            self._on_open_recent(path)

    def set_path(self, path: str):
        self.path_label.setText(f"Выбранный файл: {path}" if path else "Файл не выбран")

    def set_recent_files(self, paths):
        self.recent_list.clear()
        for p in paths:
            item = QListWidgetItem(os.path.basename(p))
            item.setData(Qt.UserRole, p)
            item.setToolTip(p)
            self.recent_list.addItem(item)


class TableScreen(QWidget):
    def __init__(self, on_back):
        super().__init__()
        self._on_back = on_back

        layout = QVBoxLayout(self)

        top = QHBoxLayout()
        self.file_label = QLabel("Файл не выбран")
        top.addWidget(self.file_label)
        top.addStretch()

        self.back_btn = QPushButton("Назад")
        self.back_btn.clicked.connect(self._on_back)
        top.addWidget(self.back_btn)

        layout.addLayout(top)

        self.table_view = QTableView()
        self.model = CsvTableModel()
        self.table_view.setModel(self.model)
        self.table_view.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        layout.addWidget(self.table_view)

    def show_table(self, path, rows, headers):
        self.file_label.setText(f"Файл: {path}")
        self.model.set_data(rows, headers)


class TextScreen(QWidget):
    def __init__(self, on_back):
        super().__init__()
        self._on_back = on_back

        layout = QVBoxLayout(self)

        top = QHBoxLayout()
        self.file_label = QLabel("Файл не выбран")
        top.addWidget(self.file_label)
        top.addStretch()

        self.back_btn = QPushButton("Назад")
        self.back_btn.clicked.connect(self._on_back)
        top.addWidget(self.back_btn)

        layout.addLayout(top)

        self.text_view = QTextEdit()
        self.text_view.setReadOnly(True)
        layout.addWidget(self.text_view)

    def show_text(self, path, text):
        self.file_label.setText(f"Файл: {path}")
        self.text_view.setPlainText(text)


class FileViewerWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Просмотрщик файлов")
        self.resize(900, 650)

        self.recent_files = []

        self.stack = QStackedWidget()
        self.setCentralWidget(self.stack)

        self.main_screen = MainScreen(
            on_select_file=self._on_select_file,
            on_open_recent=self._open_file,
        )
        self.table_screen = TableScreen(on_back=self._go_home)
        self.text_screen = TextScreen(on_back=self._go_home)

        self.stack.addWidget(self.main_screen)
        self.stack.addWidget(self.table_screen)
        self.stack.addWidget(self.text_screen)

        self.statusBar().showMessage("Готово")
        self._refresh_recent()

    def _go_home(self):
        self.stack.setCurrentWidget(self.main_screen)
        self.statusBar().showMessage("Главный экран")

    def _refresh_recent(self):
        self.main_screen.set_recent_files(self.recent_files)

    def _add_to_recent(self, path: str):
        path = os.path.abspath(path)
        if path in self.recent_files:
            self.recent_files.remove(path)
        self.recent_files.insert(0, path)
        self.recent_files = self.recent_files[:10]
        self._refresh_recent()

    def _on_select_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите файл",
            "",
            "Поддерживаемые файлы (*.csv *.xlsx *.txt *.docx);;Все файлы (*.*)",
        )
        if not path:
            return
        self._open_file(path)

    def _open_file(self, path: str):
        if not os.path.exists(path):
            QMessageBox.warning(self, "Ошибка", "Файл не найден")
            self._remove_from_recent(path)
            return

        ext = os.path.splitext(path)[1].lower()

        if ext == ".csv":
            self._open_csv(path)
        elif ext == ".xlsx":
            self._open_xlsx(path)
        elif ext == ".txt":
            self._open_txt(path)
        elif ext == ".docx":
            self._open_docx(path)
        else:
            QMessageBox.warning(
                self,
                "Неподдерживаемый формат",
                f"Формат «{ext or 'без расширения'}» не поддерживается.\n"
                "Поддерживаются: .csv, .xlsx, .txt, .docx",
            )
            return

    def _remove_from_recent(self, path: str):
        path = os.path.abspath(path)
        if path in self.recent_files:
            self.recent_files.remove(path)
            self._refresh_recent()

    def _open_csv(self, path: str):
        try:
            with open(path, "r", encoding="utf-8-sig", newline="") as f:
                reader = list(csv.reader(f))
        except UnicodeDecodeError:
            try:
                with open(path, "r", encoding="cp1251", newline="") as f:
                    reader = list(csv.reader(f))
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл:\n{e}")
                return
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл:\n{e}")
            return

        if not reader:
            QMessageBox.information(self, "Пустой файл", "Файл не содержит данных")
            return

        headers = reader[0]
        rows = reader[1:] if len(reader) > 1 else []

        self.table_screen.show_table(path, rows, headers)
        self.stack.setCurrentWidget(self.table_screen)
        self.main_screen.set_path(path)
        self._add_to_recent(path)
        self.statusBar().showMessage(f"Открыт CSV: {path}")

    def _open_xlsx(self, path: str):
        try:
            from openpyxl import load_workbook
        except ImportError:
            QMessageBox.critical(
                self,
                "Отсутствует библиотека",
                "Для чтения .xlsx установите openpyxl:\n\npip install openpyxl",
            )
            return

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            data = list(ws.iter_rows(values_only=True))
            wb.close()
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл:\n{e}")
            return

        if not data:
            QMessageBox.information(self, "Пустой файл", "Файл не содержит данных")
            return

        headers = [("" if c is None else str(c)) for c in data[0]]
        rows = [
            [("" if c is None else str(c)) for c in row]
            for row in data[1:]
        ]

        self.table_screen.show_table(path, rows, headers)
        self.stack.setCurrentWidget(self.table_screen)
        self.main_screen.set_path(path)
        self._add_to_recent(path)
        self.statusBar().showMessage(f"Открыт XLSX: {path}")

    def _open_txt(self, path: str):
        text = None
        for enc in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                with open(path, "r", encoding=enc) as f:
                    text = f.read()
                break
            except UnicodeDecodeError:
                continue
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл:\n{e}")
                return

        if text is None:
            QMessageBox.critical(
                self, "Ошибка",
                "Не удалось определить кодировку файла"
            )
            return

        self.text_screen.show_text(path, text)
        self.stack.setCurrentWidget(self.text_screen)
        self.main_screen.set_path(path)
        self._add_to_recent(path)
        self.statusBar().showMessage(f"Открыт TXT: {path}")

    def _open_docx(self, path: str):
        try:
            import docx
        except ImportError:
            QMessageBox.critical(
                self,
                "Отсутствует библиотека",
                "Для чтения .docx установите python-docx:\n\npip install python-docx",
            )
            return

        try:
            document = docx.Document(path)
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать файл:\n{e}")
            return

        self.text_screen.show_text(path, text)
        self.stack.setCurrentWidget(self.text_screen)
        self.main_screen.set_path(path)
        self._add_to_recent(path)
        self.statusBar().showMessage(f"Открыт DOCX: {path}")


def main():
    app = QApplication(sys.argv)
    window = FileViewerWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
